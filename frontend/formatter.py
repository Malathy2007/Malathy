from docx import Document
from fpdf import FPDF

def format_txt(text):
    return text

def format_docx(text):
    file_path = "legal_document.docx"

    doc = Document()
    doc.add_heading("Legal Document", 0)

    for line in text.split("\n"):
        doc.add_paragraph(line)

    doc.save(file_path)
    return file_path

def format_pdf(text):
    file_path = "legal_document.pdf"

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for line in text.split("\n"):
        pdf.multi_cell(0, 10, line)

    pdf.output(file_path)
    return file_path